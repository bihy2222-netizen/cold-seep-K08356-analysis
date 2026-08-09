#!/usr/bin/env python3
from pathlib import Path
import shutil

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE = Path("/Users/catherine/Downloads/MAG-bin-tpm")
OUT = BASE / "K08356_gene_and_MAG_dual_evidence_20260708"
FIG = OUT / "figures"
XLSX = OUT / "K08356_gene_and_MAG_dual_evidence_report.xlsx"

SHEETS = [
    ("README_Method", OUT / "README_method_and_interpretation.csv"),
    ("Gene_Total_Long", OUT / "01_gene_level_K08356_TPM/K08356_gene_total_sample_long.csv"),
    ("Gene_Habitat_Summary", OUT / "01_gene_level_K08356_TPM/K08356_gene_total_habitat_summary.csv"),
    ("Gene_Kruskal", OUT / "01_gene_level_K08356_TPM/K08356_gene_total_Kruskal.csv"),
    ("Gene_Pairwise_Wilcox", OUT / "01_gene_level_K08356_TPM/K08356_gene_total_pairwise_Wilcoxon_BH.csv"),
    ("MAG_Branch_Long", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_TPM_by_sample.csv"),
    ("MAG_Total_Long", OUT / "02_MAG_host_TPM/K08356_unique_host_MAG_total_TPM_by_sample.csv"),
    ("MAG_Proportion_Long", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_proportion_by_sample.csv"),
    ("MAG_MeanPerCopy_Long", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_mean_per_copy_by_sample.csv"),
    ("MAG_Denominators", OUT / "02_MAG_host_TPM/K08356_branch_MAG_denominators.csv"),
    ("MAG_Branch_Summary", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_habitat_summary.csv"),
    ("MAG_Proportion_Summary", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_proportion_habitat_summary.csv"),
    ("MAG_MeanPerCopy_Summary", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_mean_per_copy_habitat_summary.csv"),
    ("MAG_Branch_Kruskal", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_Kruskal.csv"),
    ("MAG_Proportion_Kruskal", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_proportion_Kruskal.csv"),
    ("MAG_MeanPerCopy_Kruskal", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_mean_per_copy_Kruskal.csv"),
    ("MAG_Branch_Pairwise", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_pairwise_Wilcoxon_BH.csv"),
    ("MAG_Proportion_Pairwise", OUT / "02_MAG_host_TPM/K08356_branch_host_MAG_proportion_pairwise_Wilcoxon_BH.csv"),
    ("MAG_Total_Kruskal", OUT / "02_MAG_host_TPM/K08356_unique_host_MAG_total_Kruskal.csv"),
    ("MAG_Total_Pairwise", OUT / "02_MAG_host_TPM/K08356_unique_host_MAG_total_pairwise_Wilcoxon_BH.csv"),
]


def read_table(path):
    return pd.read_csv(path)


def main():
    with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
        for sheet, path in SHEETS:
            read_table(path).to_excel(writer, sheet_name=sheet, index=False)
        pd.DataFrame({
            "Figure": [
                "01 Gene total K08356 TPM boxplot",
                "02 Gene total K08356 TPM sample-order bar",
                "03 MAG host TPM by branch boxplots",
                "04 MAG host total TPM boxplot",
                "05 MAG host branch TPM sample-order stacked bar",
            ],
            "Interpretation": [
                "Functional gene abundance evidence",
                "Functional gene abundance by sample/station order",
                "K08356-bearing host MAG abundance evidence by branch",
                "Total K08356-bearing host MAG abundance",
                "Spatial distribution/ecological niche of branch-associated host MAGs",
            ],
        }).to_excel(writer, sheet_name="Figures", index=False)

    wb = load_workbook(XLSX)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="D0D7DE")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(ws.columns, start=1):
            max_len = max(min(len(str(c.value)) if c.value is not None else 0, 60) for c in col)
            ws.column_dimensions[get_column_letter(idx)].width = max(10, min(max_len + 2, 45))

    ws = wb["Figures"]
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 62
    for r in range(1, 120):
        ws.row_dimensions[r].height = 22
    images = [
        ("01_gene_total_K08356_TPM_boxplot.png", "A8", 0.47),
        ("03_MAG_host_TPM_by_branch_boxplots.png", "J8", 0.44),
        ("04_MAG_host_total_TPM_boxplot.png", "A34", 0.47),
        ("02_gene_total_K08356_TPM_sample_order_bar.png", "J46", 0.36),
        ("05_MAG_host_branch_TPM_sample_order_stacked_bar.png", "A66", 0.36),
    ]
    for filename, anchor, scale in images:
        img = Image(str(FIG / filename))
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)
        ws.add_image(img, anchor)
    wb.save(XLSX)

    package_dir = BASE / "checked_corrected_tree_beautify_files_20260708" / "DUAL_GENE_AND_MAG_EVIDENCE_REPORT"
    package_fig = package_dir / "figures"
    package_fig.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX, package_dir / XLSX.name)
    for path in OUT.rglob("*.csv"):
        rel = path.relative_to(OUT)
        target = package_dir / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, target)
    for path in FIG.glob("*"):
        shutil.copy2(path, package_fig / path.name)
    print(XLSX)
    print(package_dir)


if __name__ == "__main__":
    main()
