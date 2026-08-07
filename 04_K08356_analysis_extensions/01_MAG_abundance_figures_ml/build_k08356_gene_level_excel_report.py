#!/usr/bin/env python3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter


OUTDIR = Path("/Users/catherine/Downloads/MAG-bin-tpm/k08356_gene_level_contig_style_report")
FIGDIR = OUTDIR / "figures"
XLSX = OUTDIR / "K08356_gene_level_contig_style_significance_report.xlsx"

SHEETS = [
    ("README", OUTDIR / "README_method_note.csv"),
    ("Branch_Sample_Long", OUTDIR / "K08356_gene_branch_sample_long.csv"),
    ("Total_Sample_Long", OUTDIR / "K08356_gene_total_sample_long.csv"),
    ("Branch_Habitat_Summary", OUTDIR / "K08356_gene_branch_habitat_summary.csv"),
    ("Branch_Kruskal", OUTDIR / "K08356_gene_branch_Kruskal.csv"),
    ("Branch_Pairwise_Wilcox", OUTDIR / "K08356_gene_branch_pairwise_Wilcoxon_BH.csv"),
    ("Total_Kruskal", OUTDIR / "K08356_gene_total_Kruskal.csv"),
]


def read_csv(path):
    return pd.read_csv(path)


def main():
    with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
        for sheet, path in SHEETS:
            read_csv(path).to_excel(writer, sheet_name=sheet, index=False)
        pd.DataFrame({
            "Figure": [
                "K08356 gene total TPM boxplot",
                "K08356 gene TPM by Tree11 branch boxplots",
            ],
            "File": [
                str(FIGDIR / "K08356_gene_total_TPM_boxplot.png"),
                str(FIGDIR / "K08356_gene_TPM_by_branch_boxplots.png"),
            ],
        }).to_excel(writer, sheet_name="Figures", index=False)

    wb = load_workbook(XLSX)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="D0D7DE")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(ws.columns, start=1):
            max_len = max(min(len(str(c.value)) if c.value is not None else 0, 60) for c in col)
            ws.column_dimensions[get_column_letter(idx)].width = max(10, min(max_len + 2, 42))

    ws = wb["Figures"]
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 65
    for row in range(1, 80):
        ws.row_dimensions[row].height = 22
    images = [
        ("K08356_gene_total_TPM_boxplot.png", "A5", 0.55),
        ("K08356_gene_TPM_by_branch_boxplots.png", "J5", 0.55),
    ]
    for filename, anchor, scale in images:
        img = Image(str(FIGDIR / filename))
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)
        ws.add_image(img, anchor)

    wb.save(XLSX)
    print(XLSX)


if __name__ == "__main__":
    main()
