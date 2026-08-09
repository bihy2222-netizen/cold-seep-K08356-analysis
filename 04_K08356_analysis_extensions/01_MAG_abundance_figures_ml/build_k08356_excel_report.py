#!/usr/bin/env python3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE = Path("/Users/catherine/Downloads/MAG-bin-tpm")
OUTDIR = BASE / "derepMAG_abundance_stats"
FIGDIR = OUTDIR / "figures"
XLSX = OUTDIR / "K08356_derepMAG_significance_report.xlsx"


SHEETS = [
    ("Sample_Habitat", OUTDIR / "sample_habitat_map.tsv"),
    ("Sample_Group_Summary", OUTDIR / "sample_group_summary.tsv"),
    ("Branch_Map", OUTDIR / "k08356_branch_feature_map.tsv"),
    ("Tree_Habitat_Audit", OUTDIR / "itol_habitat_colorstrip_audit.tsv"),
    ("Tree_Corrections", OUTDIR / "tree_habitat_corrections.tsv"),
    ("Abundance_Long", OUTDIR / "feature_abundance_by_sample.tsv"),
    ("Proportion_Long", OUTDIR / "feature_proportion_by_sample.tsv"),
    ("Kruskal_TPM", OUTDIR / "feature_kruskal.tsv"),
    ("Pairwise_TPM", OUTDIR / "feature_pairwise_wilcoxon.tsv"),
    ("Kruskal_Proportion", OUTDIR / "feature_proportion_kruskal.tsv"),
    ("Pairwise_Proportion", OUTDIR / "feature_proportion_pairwise_wilcoxon.tsv"),
    ("NMDS_ANOSIM", OUTDIR / "nmds_anosim_stats.tsv"),
]


def read_tsv(path):
    return pd.read_csv(path, sep="\t")


def add_readme(writer):
    nmds = read_tsv(OUTDIR / "nmds_anosim_stats.tsv")
    nmds_map = dict(zip(nmds["Metric"], nmds["Value"]))
    kw = read_tsv(OUTDIR / "feature_kruskal.tsv")
    total_p = kw.loc[kw["Feature"] == "K08356_bearing_unique_MAG_total_TPM", "Kruskal_p"].iloc[0]
    rows = [
        ["Item", "Value"],
        ["Input matrix", str(BASE / "merged_tpm_matrix.txt")],
        ["Branch source workbook", "/Users/catherine/Downloads/aioA 蛋白序列建树/MAG 补充 idra-tree/latest_MAG_tree_beautify_files/MAG_tree_clade_statistics_like_tree11.xlsx"],
        ["MAG branch records", 48],
        ["Unique K08356-bearing MAGs", 47],
        ["Habitat rule", "SY*=IS; S1/S2/S4/SQ*=AS; C*/ES_2/S13/S14/S15=ES; S3/NS/R2111*=NS"],
        ["Tree habitat correction", "The source tree habitat colorstrip had stale labels for S13/S15; corrected copies and audit tables are included."],
        ["Total K08356-bearing unique MAG Kruskal P", total_p],
        ["ANOSIM R", nmds_map.get("ANOSIM_R")],
        ["ANOSIM P", nmds_map.get("ANOSIM_P")],
        ["NMDS stress", nmds_map.get("NMDS_stress")],
        ["Notes", "Boxplots use log10(TPM+1) for abundance; proportion plots use branch abundance / unique K08356-bearing MAG total."],
        ["Notes", "Pairwise tests are Wilcoxon rank-sum / Mann-Whitney U with BH-adjusted P values."],
    ]
    pd.DataFrame(rows[1:], columns=rows[0]).to_excel(writer, sheet_name="README", index=False)


def write_workbook():
    with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
        add_readme(writer)
        for sheet, path in SHEETS:
            read_tsv(path).to_excel(writer, sheet_name=sheet, index=False)
        pd.DataFrame({"Figure": [
            "K08356 total abundance boxplot",
            "K08356 branch abundance boxplots",
            "K08356 branch proportion boxplots",
            "K08356 branch composition NMDS + ANOSIM",
        ]}).to_excel(writer, sheet_name="Figures", index=False)


def format_workbook():
    wb = load_workbook(XLSX)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    title_fill = PatternFill("solid", fgColor="1F4E78")
    title_font = Font(color="FFFFFF", bold=True)
    header_font = Font(bold=True, color="1F1F1F")
    thin = Side(style="thin", color="D0D7DE")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        ws.auto_filter.ref = ws.dimensions
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value), 60))
            ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(max_len + 2, 42))

    ws = wb["README"]
    ws.insert_rows(1, 2)
    ws["A1"] = "K08356 derepMAG abundance significance report"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.merge_cells("A1:B1")
    ws["A2"] = "Generated from merged_tpm_matrix.txt and MAG_tree_clade_statistics_like_tree11.xlsx"
    ws["A2"].font = Font(italic=True, color="555555")
    ws.merge_cells("A2:B2")
    ws.freeze_panes = "A5"

    ws_fig = wb["Figures"]
    ws_fig.sheet_view.showGridLines = False
    for row in range(2, 80):
        ws_fig.row_dimensions[row].height = 22
    ws_fig.column_dimensions["A"].width = 38
    ws_fig.column_dimensions["B"].width = 3
    ws_fig.column_dimensions["C"].width = 38

    figures = [
        ("K08356_total_abundance_boxplot.png", "A3"),
        ("K08356_branch_abundance_boxplots.png", "A28"),
        ("K08356_branch_proportion_boxplots.png", "J3"),
        ("K08356_branch_NMDS_ANOSIM.png", "J35"),
    ]
    for filename, anchor in figures:
        path = FIGDIR / filename
        img = Image(str(path))
        img.width = int(img.width * 0.62)
        img.height = int(img.height * 0.62)
        ws_fig.add_image(img, anchor)

    wb.save(XLSX)


def main():
    write_workbook()
    format_workbook()
    print(XLSX)


if __name__ == "__main__":
    main()
